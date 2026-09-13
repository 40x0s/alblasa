#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""معاينة هندسية للـ PPTX (PIL): مربعات + نص ملفوف بنفس العرض. ليست معاينة شكل عربي —
هدفها كشف التراكب والخروج خارج الشريحة فقط."""
import sys, os
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

DPI = 100
F = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
_cache = {}
def font(name, size_pt):
    key = (name, int(size_pt))
    if key not in _cache:
        fn = os.path.join(F, "DejaVuSansMono.ttf" if name == "Consolas" else "DejaVuSans.ttf")
        _cache[key] = ImageFont.truetype(fn, max(6, int(size_pt * DPI / 72)))
    return _cache[key]

def w_of(txt, f, arabic):
    # العربية أضيق من اللاتينية في Segoe UI؛ معامل تقريبي للعرض
    return f.getlength(txt) * (0.665 if arabic else 1.0)

def has_ar(t):
    return any('\u0600' <= c <= '\u06ff' for c in t)

def render(path, outdir, only=None):
    prs = Presentation(path)
    SW = int(prs.slide_width / 914400 * DPI)
    SH = int(prs.slide_height / 914400 * DPI)
    os.makedirs(outdir, exist_ok=True)
    reports = []
    for idx, sl in enumerate(prs.slides, 1):
        if only and idx not in only:
            continue
        im = Image.new("RGB", (SW, SH), "white")
        dr = ImageDraw.Draw(im)
        for sh in sl.shapes:
            x = int(sh.left / 914400 * DPI); y = int(sh.top / 914400 * DPI)
            w = int(sh.width / 914400 * DPI); h = int(sh.height / 914400 * DPI)
            fill = None
            try:
                if sh.fill.type is not None and str(sh.fill.type) == "MSO_FILL_TYPE.SOLID":
                    pass
            except Exception:
                pass
            # لون التعبئة (تقريبي عبر XML)
            try:
                srgb = sh.fill.fore_color.rgb
                fill = (srgb[0], srgb[1], srgb[2])
            except Exception:
                fill = None
            if fill is not None:
                dr.rectangle([x, y, x + w, y + h], fill=fill)
                ln = None
                try:
                    lc = sh.line.color.rgb
                    ln = (lc[0], lc[1], lc[2])
                except Exception:
                    ln = None
                if ln:
                    dr.rectangle([x, y, x + w, y + h], outline=ln, width=2)
            if not sh.has_text_frame:
                continue
            if x < 0 or y < 0 or x + w > SW or y + h > SH:
                reports.append((idx, "OFF-SLIDE", x, y, w, h, sh.text_frame.text[:30]))
            ty = y + int(0.02 * DPI)
            for p in sh.text_frame.paragraphs:
                runs = [(r.text, (r.font.name or ""), (r.font.size.pt if r.font.size else 18),
                         (r.font.color.rgb[0:3] if r.font.color and r.font.color.type else (20, 26, 34)))
                        for r in p.runs]
                if not any(t for t, *_ in runs):
                    continue
                fnt = font(runs[0][1], runs[0][2])
                words = []
                for t, nm, sz, col in runs:
                    f2 = font(nm, sz)
                    for piece in t.split(" "):
                        words.append((piece, f2, col, sz))
                gap = fnt.getlength(" ") * 0.9
                lines, cur, cw = [[]], [], 0.0
                for piece, f2, col, sz in words:
                    wp = w_of(piece, f2, has_ar(piece))
                    if cw + wp > max(w - 6, 20) and cur:
                        lines.append(cur); cur, cw = [(piece, f2, col)], wp
                    else:
                        cur.append((piece, f2, col)); cw += wp + gap
                if cur:
                    lines.append(cur)
                align = p.alignment
                rtl = p._pPr is not None and p._pPr.get("algn") in ("r", None)
                lh = max((r[2] for r in runs), default=16) * DPI / 72 * 1.18
                for ln in lines:
                    lw = sum(w_of(t, f2, has_ar(t)) for t, f2, _ in ln) + gap * max(0, len(ln) - 1)
                    cx = (x + w - lw - 4) if rtl else (x + 4)
                    for t, f2, col in ln:
                        dr.text((cx, ty), t, font=f2, fill=col)
                        cx += w_of(t, f2, has_ar(t)) + gap
                    ty += lh
                ty += runs[0][2] * DPI / 72 * 0.22
            if ty > y + h + 22:
                reports.append((idx, "OVERFLOW-BOTTOM", round(ty - (y + h)), "px",
                                sh.text_frame.text[:40].replace("\n", " "), ""))
        out = os.path.join(outdir, f"slide_{idx:02d}.png")
        im.save(out)
    for r in reports:
        print(r)
    print("rendered ->", outdir, " issues:", len(reports))

if __name__ == "__main__":
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "MizanPro_Crack_Workshop.pptx")
    only = {int(a) for a in sys.argv[1:]} or None
    render(path, "/home/user/build/ppt", only)
